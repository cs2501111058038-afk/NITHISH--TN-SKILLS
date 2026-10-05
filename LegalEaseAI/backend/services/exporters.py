import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF


ASSET_DIR = (
    Path(__file__).resolve().parents[2] / "assets"
)

LOGO_PATH = ASSET_DIR / "logo.png"


def _lines(text: str) -> list[str]:

    return [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]


def _clean_heading(line: str) -> str:

    return re.sub(
        r"^#+\s*",
        "",
        line,
    ).strip()


def format_txt(text: str) -> bytes:

    return text.strip().encode("utf-8")


def format_docx(
    text: str,
    doc_type: str,
    logo_path: Path | None = None,
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles

    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    logo = logo_path or LOGO_PATH

    if logo.exists():

        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run()

        run.add_picture(
            str(logo),
            width=Inches(1.35),
        )

    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(
        doc_type.upper()
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    in_terms = False

    terms_for_table = []

    for line in _lines(text):

        # Skip Markdown document title because we already add our own.
        if line.startswith("# ") and not line.startswith("## "):
            continue

        # Section heading
        if line.startswith("## "):

            heading = _clean_heading(line)

            in_terms = (
                "TERMS AND CONDITIONS"
                in heading.upper()
            )

            paragraph = document.add_paragraph()

            run = paragraph.add_run(heading)

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

            continue

        # Numbered clause
        numbered = re.match(
            r"^\d+\.\s+(.+)$",
            line,
        )

        if numbered:

            clause = numbered.group(1)

            if in_terms:

                terms_for_table.append(clause)

            else:

                paragraph = document.add_paragraph(
                    style="List Number"
                )

                paragraph.add_run(clause)

            continue

        # Bullet
        if line.startswith(("- ", "* ")):

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                line[2:]
            )

            continue

        document.add_paragraph(line)

    # Terms table
    if terms_for_table:

        document.add_paragraph()

        table = document.add_table(
            rows=1,
            cols=2,
        )

        table.style = "Table Grid"

        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term / Condition"

        for index, clause in enumerate(
            terms_for_table,
            start=1,
        ):

            cells = table.add_row().cells

            cells[0].text = str(index)
            cells[1].text = clause

    # Footer
    footer = section.footer.paragraphs[0]

    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer.add_run(
        "LegalEase - AI-assisted draft. "
        "Review with a qualified legal professional."
    )

    footer_run.font.size = Pt(8)

    output = BytesIO()

    document.save(output)

    return output.getvalue()


class LegalEasePDF(FPDF):

    def __init__(self, doc_type: str):

        super().__init__()

        self.doc_type = doc_type

        self.set_auto_page_break(
            auto=True,
            margin=18,
        )

    def header(self):

        if LOGO_PATH.exists():

            try:

                self.image(
                    str(LOGO_PATH),
                    x=95,
                    y=8,
                    w=20,
                )

            except Exception:
                pass

        self.set_y(30)

        self.set_font(
            "Helvetica",
            "B",
            13,
        )

        self.cell(
            0,
            8,
            self.doc_type.upper(),
            align="C",
        )

        self.ln(10)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            size=7,
        )

        self.cell(
            0,
            8,
            "LegalEase - AI-assisted draft. "
            "Review with a qualified legal professional.",
            align="C",
        )


def format_pdf(
    text: str,
    doc_type: str,
) -> bytes:

    pdf = LegalEasePDF(doc_type)

    pdf.add_page()

    for line in _lines(text):

        if line.startswith("# ") and not line.startswith("## "):
            continue

        if line.startswith("## "):

            pdf.set_font(
                "Helvetica",
                "B",
                11,
            )

            pdf.ln(2)

            pdf.multi_cell(
                0,
                7,
                _clean_heading(line),
            )

            pdf.ln(1)

            continue

        if line.startswith(("- ", "* ")):

            pdf.set_font(
                "Helvetica",
                size=10,
            )

            pdf.multi_cell(
                0,
                6,
                "- " + line[2:],
            )

            continue

        numbered = re.match(
            r"^\d+\.\s+(.+)$",
            line,
        )

        if numbered:

            pdf.set_font(
                "Helvetica",
                size=10,
            )

            pdf.multi_cell(
                0,
                6,
                line,
            )

            continue

        pdf.set_font(
            "Helvetica",
            size=10,
        )

    return bytes(
        pdf.output()
    )
        

def create_txt(text: str) -> bytes:
    return format_txt(text)


def create_docx(
    text: str,
    doc_type: str,
    logo_path: Path | None = None,
) -> bytes:
    return format_docx(
        text=text,
        doc_type=doc_type,
        logo_path=logo_path,
    )


def create_pdf(
    text: str,
    doc_type: str,
) -> bytes:
    return format_pdf(
        text=text,
        doc_type=doc_type,
    )